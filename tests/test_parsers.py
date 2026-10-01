"""文件格式扩展测试:校验白名单、解析分派、各格式解析往返。

每种格式用对应库真实生成临时文件再解析;表格类(XLSX/CSV)行级成块且注入表头。

无需 pytest，可直接运行::

    uv run python tests/test_parsers.py
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

# Windows 控制台默认 GBK,统一按 UTF-8 输出
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# 保证可从项目根导入 app 包
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def test_validate_file_extension_accepts_new_types():
    """扩展名白名单覆盖新格式,拒绝未知类型。"""

    def scenario() -> None:
        from app.utils.file import validate_file_extension

        for name in ("a.docx", "b.xlsx", "c.pptx", "d.md", "e.csv", "f.html", "g.png", "h.jpeg", "i.txt", "j.pdf"):
            assert validate_file_extension(name), f"应接受 {name}"
        for name in ("x.exe", "y.xls", "z.doc", "w.ppt"):
            assert not validate_file_extension(name), f"应拒绝 {name}"

    scenario()


def test_get_parser_dispatch():
    """get_parser 按扩展名分派;图片需 Vision 配置;未知类型报错。"""

    def scenario() -> None:
        from app.rag.parser.dispatch import get_parser
        from app.rag.parser.docx_parser import DocxParser
        from app.rag.parser.image import ImageParser
        from app.rag.parser.pdf import PdfParser
        from app.rag.parser.txt import TxtParser

        assert isinstance(get_parser("txt"), TxtParser)
        assert isinstance(get_parser("md"), TxtParser), "Markdown 复用 TxtParser"
        assert isinstance(get_parser("pdf"), PdfParser)
        assert isinstance(get_parser("docx"), DocxParser)
        assert isinstance(get_parser("png", vision_config={"base_url": "x", "api_key": "y", "model": "z"}), ImageParser)
        try:
            get_parser("png")
        except ValueError as e:
            assert "Vision" in str(e)
        else:
            raise AssertionError("未配置 Vision 应报错")
        try:
            get_parser("xyz")
        except ValueError:
            pass
        else:
            raise AssertionError("未知类型应报错")

    scenario()


def test_docx_parser_sections_and_tables():
    """DOCX:标题样式分节,表格成块。"""

    def scenario() -> None:
        import docx

        from app.rag.parser.docx_parser import DocxParser

        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "t.docx"
            d = docx.Document()
            d.add_heading("第一节", level=1)
            d.add_paragraph("第一段正文。")
            d.add_paragraph("第二段正文。")
            d.add_heading("第二节", level=1)
            d.add_paragraph("第二段落。")
            table = d.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = "姓名"
            table.rows[0].cells[1].text = "年龄"
            table.rows[1].cells[0].text = "张三"
            table.rows[1].cells[1].text = "20"
            d.save(str(p))

            import asyncio
            blocks = asyncio.run(DocxParser().parse(p))
            contents = [b.content for b in blocks]
            assert any("第一段正文" in c and "第二段正文" in c for c in contents), f"节内段落应聚合,实际 {contents}"
            assert any("张三" in c and "姓名" in c for c in contents), f"表格应成块且带表头,实际 {contents}"
            titles = {b.section_title for b in blocks}
            assert "第一节" in titles and "第二节" in titles

    scenario()


def test_xlsx_parser_row_blocks():
    """XLSX:每行一块,表头注入行文本。"""

    def scenario() -> None:
        import asyncio

        import openpyxl

        from app.rag.parser.xlsx_parser import XlsxParser

        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "t.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "名单"
            ws.append(["姓名", "年龄"])
            ws.append(["张三", 20])
            ws.append(["李四", 21])
            wb.save(str(p))

            blocks = asyncio.run(XlsxParser().parse(p))
            assert len(blocks) == 2, f"应每行一块,实际 {len(blocks)}"
            assert "姓名: 张三" in blocks[0].content and "年龄: 20" in blocks[0].content, blocks[0].content
            assert blocks[0].section_title == "名单"

    scenario()


def test_pptx_parser_slide_blocks():
    """PPTX:每页一块,页标题进 section_title。"""

    def scenario() -> None:
        import asyncio

        from pptx import Presentation

        from app.rag.parser.pptx_parser import PptxParser

        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "t.pptx"
            prs = Presentation()
            s1 = prs.slides.add_slide(prs.slide_layouts[0])
            s1.shapes.title.text = "课件第一章"
            s1.placeholders[1].text = "本章介绍基础概念。"
            s2 = prs.slides.add_slide(prs.slide_layouts[0])
            s2.shapes.title.text = "课件第二章"
            s2.placeholders[1].text = "本章介绍进阶内容。"
            prs.save(str(p))

            blocks = asyncio.run(PptxParser().parse(p))
            assert len(blocks) == 2, f"应每页一块,实际 {len(blocks)}"
            assert blocks[0].section_title == "课件第一章"
            assert "基础概念" in blocks[0].content
            assert blocks[1].section_title == "课件第二章"

    scenario()


def test_csv_parser_row_blocks_and_gbk():
    """CSV:行级成块;GB18030 编码容错。"""

    def scenario() -> None:
        import asyncio

        from app.rag.parser.csv_parser import CsvParser

        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "t.csv"
            p.write_bytes("姓名,年龄\n张三,20\n李四,21\n".encode("gb18030"))

            blocks = asyncio.run(CsvParser().parse(p))
            assert len(blocks) == 2, f"应每行一块,实际 {len(blocks)}"
            assert "姓名: 张三" in blocks[0].content and "年龄: 20" in blocks[0].content

    scenario()


def test_html_parser_strips_script_and_sections():
    """HTML:script/style 不入正文;h1-h6 分节。"""

    def scenario() -> None:
        import asyncio

        from app.rag.parser.html_parser import HtmlParser

        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "t.html"
            p.write_text(
                "<html><head><title>通知</title><script>var x=1;</script>"
                "<style>.a{}</style></head><body>"
                "<h1>开学通知</h1><p>九月一日开学。</p>"
                "<h2>注意事项</h2><p>携带学生证。</p>"
                "</body></html>",
                encoding="utf-8",
            )
            blocks = asyncio.run(HtmlParser().parse(p))
            joined = "".join(b.content or "" for b in blocks)
            assert "九月一日开学" in joined and "携带学生证" in joined
            assert "var x=1" not in joined and ".a{}" not in joined, "script/style 不得混入正文"
            titles = {b.section_title for b in blocks}
            assert "开学通知" in titles and "注意事项" in titles

    scenario()


def test_chunker_keeps_table_blocks():
    """切块器放行 TABLE 块(表格行块不得被丢弃)。"""

    def scenario() -> None:
        from app.rag.chunker.text_chunker import chunk_blocks
        from app.rag.models.blocks import BlockType, ContentBlock

        blocks = [ContentBlock(block_type=BlockType.TABLE, content="姓名: 张三 | 年龄: 20", source_index=0)]
        chunks = chunk_blocks(blocks, document_id=1, knowledge_base_id=1, template="section")
        assert len(chunks) == 1 and "张三" in chunks[0].content, f"TABLE 块应参与切块,实际 {chunks}"

    scenario()


def main() -> int:
    tests = [
        obj
        for name, obj in sorted(globals().items())
        if name.startswith("test_") and callable(obj)
    ]
    failed = 0
    for test in tests:
        try:
            test()
        except Exception as e:
            failed += 1
            print(f"FAIL {test.__name__}: {type(e).__name__}: {e}")
        else:
            print(f"PASS {test.__name__}")
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
