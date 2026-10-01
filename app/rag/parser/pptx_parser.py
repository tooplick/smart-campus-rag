from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from app.rag.models.blocks import BlockType, ContentBlock
from app.rag.parser.base import BaseParser


class PptxParser(BaseParser):
    """PPTX 解析:每页一块,页标题进 section_title,页内文本与表格行并入正文。"""

    async def parse(self, file_path: Path) -> list[ContentBlock]:
        presentation = Presentation(str(file_path))
        blocks: list[ContentBlock] = []

        for index, slide in enumerate(presentation.slides):
            title_shape = slide.shapes.title
            title_id = title_shape.shape_id if title_shape is not None else None
            title = title_shape.text.strip() if title_shape is not None and title_shape.text else None

            lines: list[str] = []
            for shape in slide.shapes:
                if shape.shape_id == title_id:
                    continue
                if shape.has_text_frame and shape.text_frame.text.strip():
                    lines.append(shape.text_frame.text.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        if any(cells):
                            lines.append(" | ".join(cells))

            # 页标题并入正文(title+body):只进 section_title 会丢检索词
            if title:
                lines.insert(0, title)
            content = "\n".join(lines).strip()
            if content or title:
                blocks.append(ContentBlock(
                    block_type=BlockType.TEXT,
                    content=content or title,
                    page_number=index + 1,
                    section_title=title,
                    source_index=len(blocks),
                ))
        return blocks
