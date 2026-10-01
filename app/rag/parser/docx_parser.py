from __future__ import annotations

from pathlib import Path

import docx

from app.rag.models.blocks import BlockType, ContentBlock
from app.rag.parser.base import BaseParser
from app.rag.parser.tabular import rows_to_blocks


class DocxParser(BaseParser):
    """DOCX 解析:标题样式分节聚合段落,表格按行成块(表头注入)。"""

    async def parse(self, file_path: Path) -> list[ContentBlock]:
        document = docx.Document(str(file_path))
        blocks: list[ContentBlock] = []
        current_title: str | None = None
        buf: list[str] = []

        def flush() -> None:
            nonlocal buf
            content = "\n".join(buf).strip()
            if content:
                blocks.append(ContentBlock(
                    block_type=BlockType.TEXT,
                    content=content,
                    section_title=current_title,
                    source_index=len(blocks),
                ))
            buf = []

        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style_name = (para.style.name or "").lower()
            # Word 标题样式(英/中界面命名均覆盖)
            if style_name.startswith("heading") or style_name.startswith("标题"):
                flush()
                current_title = text
            # 标题段也保留进正文(title+body):只进 section_title 会丢检索词
            buf.append(para.text)
        flush()

        # 表格按行成块;表格出现在正文末尾的段落之后(简化:不重排 body 流)
        for table in document.tables:
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            blocks.extend(rows_to_blocks(rows, section_title=current_title, source_start=len(blocks)))
        return blocks
